"""生成可直接打开的空间框架例题；数值由真实工具返回，输出默认写入 results。"""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT), str(ROOT / "src")]

from agent import Session  # noqa: E402
import capsule  # noqa: E402
from sections import i_section  # noqa: E402


def generate(output_dir: Path) -> dict:
    output_dir = output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    capsule.DEFAULT_CAPSULE_DIR = output_dir / "capsules"
    os.chdir(output_dir)
    session = Session()
    calls = []

    def call(name, **arguments):
        started = time.perf_counter()
        result = session.dispatch(name, arguments)
        calls.append({"tool": name, "arguments": arguments, "ok": result.ok,
                      "payload": result.payload,
                      "elapsed_s": round(time.perf_counter() - started, 4)})
        if not result.ok:
            raise RuntimeError(f"{name}: {result.payload}")
        return result.payload

    call("define_materials_and_sections", materials=[{
        "name": "Q355", "E": 206e9, "nu": 0.3, "density": 7850,
        "yield_stress": 355e6, "allow_tension": 305e6, "allow_compression": 305e6,
    }], sections=[i_section("COL-H350x350x12x19", .350, .350, .012, .019),
                   i_section("BEAM-H350x175x7x11", .350, .175, .007, .011)])
    call("generate_frame", spans=[6, 6], storeys=[3.6], bays=[4],
         column_section="COL-H350x350x12x19", beam_section="BEAM-H350x175x7x11",
         material="Q355", base="fixed")
    nodes = {node["id"]: node for node in session.model["nodes"]}
    beam_ids = [member["id"] for member in session.model["members"]
                if nodes[member["i"]]["z"] == nodes[member["j"]]["z"] == 3.6]
    roof_ids = [node["id"] for node in nodes.values() if node["z"] == 3.6]
    cases = [{"name": name, "member_loads": [
        {"name": f"{name}-M{mid}", "member": mid, "w": [0, 0, -load]}
        for mid in beam_ids]} for name, load in [("DL", 5000), ("LL", 3000)]]
    cases.append({"name": "WX", "nodal_loads": [
        {"name": f"WX-N{nid}", "node": nid, "load": [2000, 0, 0, 0, 0, 0]}
        for nid in roof_ids]})
    combos = [{"name": "ULS-G", "factors": {"DL": 1.3, "LL": 1.5}},
              {"name": "ULS-GW", "factors": {"DL": 1.3, "LL": 1.5, "WX": 1.5}},
              {"name": "SLS", "factors": {"DL": 1, "LL": 1}}]
    call("set_load_cases", cases=cases, combos=combos)
    call("validate_model")
    call("solve_model")
    reactions = {}
    for case in ("DL", "LL", "WX", "ULS-G", "ULS-GW", "SLS"):
        reactions[case] = call("query_results", what="reactions", case=case)
    # 独立从题设长度和外力求合力，不从求解器读荷载合计。
    expected = {"DL": [0, 0, 180], "LL": [0, 0, 108], "WX": [-12, 0, 0],
                "ULS-G": [0, 0, 396], "ULS-GW": [-18, 0, 396], "SLS": [0, 0, 288]}
    import numpy as np
    for case, target in expected.items():
        result = session.solution[case]
        actual = sum((result.R[session.frame.node_dofs(nid)][:3]
                      for nid in session.frame.supports), np.zeros(3)) / 1000
        if not np.allclose(actual, target, atol=1e-8, rtol=1e-9):
            raise AssertionError((case, actual.tolist(), target))
    displacement = call("query_results", what="max_displacement", case="ULS-GW")
    deflection = call("query_results", what="max_deflection", case="SLS", span=6)
    strength = call("check_strength")
    envelope = call("query_envelope", component="Mz")
    modal = call("modal_analysis", num_modes=6)
    report = call("write_report", case="ULS-GW", fmt="both", filename="双跨空间钢框架计算书")
    # 本例交付计算书的表头与首行同页，避免只在页尾留下表头。
    from docx import Document
    from docx.oxml import OxmlElement
    document = Document(report["files"]["docx"])
    for table in document.tables:
        header = table.rows[0]
        header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell in header.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    document.save(report["files"]["docx"])
    model_path = output_dir / "双跨空间钢框架.json"
    model_path.write_text(json.dumps(session.model, ensure_ascii=False, indent=2), encoding="utf-8")
    record = {"description": "双跨单层空间钢框架教学例题；组合系数及允许应力为演示设定。",
              "model": str(model_path), "nodes": len(nodes), "members": len(session.model["members"]),
              "supports": len(session.model["supports"]), "roof_beams": beam_ids,
              "roof_nodes": roof_ids, "independent_reaction_check_kN": expected,
              "displacement": displacement, "deflection": deflection,
              "strength": strength, "envelope": envelope, "modal": modal,
              "report": report, "calls": calls}
    (output_dir / "工具结果.json").write_text(
        json.dumps(record, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps({key: record[key] for key in ("model", "nodes", "members", "supports",
                     "displacement", "deflection")}, ensure_ascii=False, indent=2))
    return record


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "results/defense_space_frame")
    generate(parser.parse_args().output)
