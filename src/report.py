"""分析报告。

把一次分析的全部结论汇总成一份可交付的报告：模型概况、荷载、校验、
结果表、内力图、控制组合，以及（算过的话）模态与屈曲。

**每个数字都来自工具返回**，和界面、和评测集是同一条链路——报告里不会
出现任何"另算一遍"的数。这一条是整个项目的底线，报告只是把它们排版。

图是现成的静态出图函数产的 PNG，与界面上的交互图共用 `viz_theme` 配色，
所以报告和屏幕不会各说各话。

产出两种格式：Markdown（轻，便于复制粘贴和版本管理）和 docx（交作业用）。
两者的内容由同一个 `gather()` 生成，不会漂。
"""

from __future__ import annotations

import datetime as _dt
from pathlib import Path
from typing import Any

_COMPONENT_LABEL = {"N": "轴力 N", "Vy": "剪力 Vy", "Vz": "剪力 Vz",
                    "T": "扭矩 T", "My": "弯矩 My", "Mz": "弯矩 Mz"}


def _safe(name: str) -> str:
    return "".join(c if c.isalnum() or c in "-_" else "_" for c in str(name))


def _describe(path: Path) -> dict[str, Any]:
    """图片路径 + 实际像素尺寸。

    尺寸要带给排版层：裁掉留白之后图片接近正方形，还按写死的 560×380 嵌入
    会把图拉变形。让排版按真实宽高比缩放。
    """
    out: dict[str, Any] = {"path": str(path), "w": 0, "h": 0}
    try:
        from PIL import Image
        with Image.open(path) as im:
            out["w"], out["h"] = im.size
    except Exception:                       # noqa: BLE001
        pass
    return out


def _trim(path: Path, pad: int = 8) -> Path:
    """裁掉图片四周的纯白边。

    matplotlib 的三维图周围有一大圈留白，直接塞进报告会变成"大白框里一张小图"。
    裁完再嵌，同样的版面里图能大出一大截。裁不动就原样返回——
    修图失败不该让整份报告出不来。
    """
    try:
        from PIL import Image, ImageChops
    except ImportError:
        return path
    try:
        im = Image.open(path).convert("RGB")
        bg = Image.new("RGB", im.size, im.getpixel((0, 0)))
        box = ImageChops.difference(im, bg).getbbox()
        if box is None:
            return path
        left, top, right, bottom = box
        im.crop((max(0, left - pad), max(0, top - pad),
                 min(im.width, right + pad), min(im.height, bottom + pad))).save(path)
    except Exception:                       # noqa: BLE001
        return path
    return path


def gather(session, case: str | None = None, out_dir: Path | None = None,
           figures: bool = True) -> dict[str, Any]:
    """收集报告需要的一切。session 是已经求解过的 agent.Session。

    模态与屈曲**不强制**：没有密度就不做模态，全受拉就不做屈曲，
    对应小节直接略过并说明原因——比塞一段"分析失败"要干净。
    """
    if session.solution is None:
        raise ValueError("还没有求解结果，先调用 solve_model")

    name = case or session._controlling_case()
    out_dir = Path(out_dir or "results")
    out_dir.mkdir(parents=True, exist_ok=True)
    U = session.units
    model = session.model
    frame = session.frame

    doc: dict[str, Any] = {
        "generated": _dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "case": name,
        "units": {"model": model.get("units", "N-m-Pa"),
                  "report": f"位移 {U.disp_unit}、力 {U.force_unit}、"
                            f"弯矩 {U.moment_unit}"},
        "model": {
            "nodes": len(model["nodes"]),
            "members": len(model["members"]),
            "analysis_nodes": len(frame.nodes),
            "analysis_elements": len(frame.members),
            "supports": len(frame.supports),
            "fully_fixed": sum(1 for m in frame.supports.values() if all(m)),
            "materials": [{"名称": m["name"], "E": m["E"], "ν": m["nu"],
                           "密度": m.get("density", 0.0)}
                          for m in model["materials"]],
            "sections": [{"名称": s["name"], "A": s["A"], "Iy": s["Iy"],
                          "Iz": s["Iz"], "J": s["J"]}
                         for s in model["sections"]],
        },
        "cases": [c.get("name") for c in model.get("load_cases", [])],
        "combos": [c.get("name") for c in model.get("combos", [])],
        "figures": {},
        "skipped": {},
    }

    # --- 求解与校验 ---
    solved = session.solve_model()
    doc["checks"] = [
        {"工况 / 组合": k,
         "最大节点位移 (mm)": v["max_displacement_mm"],
         "静力平衡": "通过" if v["equilibrium_ok"] else "未通过",
         "平衡残差": f"{v['equilibrium_residual']:.2e}"}
        for k, v in solved.payload["cases"].items()]

    # --- 位移 ---
    nodal = session.query_results(what="max_displacement", case=name).payload
    inner = session.query_results(what="max_deflection", case=name).payload
    doc["displacement"] = {
        "node": nodal["node"],
        "coordinates": nodal["coordinates_xyz"],
        "nodal_mm": nodal["magnitude_mm"],
        "deflection_mm": inner["magnitude_mm"],
        "at_member": inner["at_member"],
        "at_x_m": inner["at_x_m"],
        "member_length_m": inner["member_length_m"],
        "length_over_deflection": inner["member_length_over_deflection"],
    }

    # --- 反力 ---
    reac = session.query_results(what="reactions", case=name).payload
    doc["reactions"] = {
        "total_vertical": reac["vertical_total_kN"],
        "rows": [{"节点": int(nid),
                  "x": v["xyz"][0], "y": v["xyz"][1], "z": v["xyz"][2],
                  "Fx": v["R"][0], "Fy": v["R"][1], "Fz": v["R"][2]}
                 for nid, v in reac["reactions"].items()],
    }

    # --- 内力极值与控制组合 ---
    rows = []
    for comp in ("N", "Vy", "Vz", "T", "My", "Mz"):
        got = session.query_envelope(component=comp)
        if not got.ok:
            continue
        p = got.payload
        rows.append({"分量": _COMPONENT_LABEL[comp], "最不利值": p["peak"],
                     "单位": p["unit"], "杆件": p["at_member"],
                     "位置 x (m)": p["at_x_m"], "控制组合": p["governing_case"]})
    doc["envelope"] = rows
    if rows:
        first = session.query_envelope(component="Mz").payload
        doc["never_governs"] = first.get("never_governs_anywhere", [])

    # --- 模态 ---
    modal = session.modal_analysis(num_modes=6)
    if modal.ok:
        doc["modal"] = modal.payload
    else:
        doc["skipped"]["模态分析"] = modal.payload.get(
            "error", str(modal.payload.get("errors", "")))[:120]

    # --- 屈曲 ---
    buck = session.buckling_analysis(case=name, num_modes=3)
    if buck.ok:
        doc["buckling"] = buck.payload
    else:
        doc["skipped"]["屈曲分析"] = buck.payload.get(
            "error", str(buck.payload.get("errors", "")))[:120]

    # --- 图 ---
    if figures:
        try:
            from plot3d import plot_axial, plot_deformed, plot_diagram
        except ImportError as exc:
            doc["skipped"]["图形"] = f"绘图依赖缺失：{exc}"
        else:
            tag = _safe(name)
            result_view = session.result_db.solution_view(frame)
            plans = [("变形图", plot_deformed, out_dir / f"deformed_{tag}.png", ()),
                     ("轴力图", plot_axial, out_dir / f"axial_{tag}.png", ()),
                     ("弯矩图 Mz", plot_diagram, out_dir / f"Mz_{tag}.png", ("Mz",)),
                     ("剪力图 Vy", plot_diagram, out_dir / f"Vy_{tag}.png", ("Vy",))]
            for label, fn, path, extra in plans:
                try:
                    if extra:
                        fn(frame, result_view, extra[0], name, path,
                           mapping=session.compilation.mapping)
                    else:
                        fn(frame, result_view, name, path,
                           mapping=session.compilation.mapping)
                    doc["figures"][label] = _describe(_trim(path))
                except Exception as exc:          # noqa: BLE001 出图失败不该毁掉报告
                    doc["skipped"][label] = f"{type(exc).__name__}: {exc}"
    return doc


# --------------------------------------------------------------- Markdown

def _table(rows: list[dict], columns: list[str] | None = None) -> list[str]:
    if not rows:
        return ["（无）", ""]
    cols = columns or list(rows[0])
    out = ["| " + " | ".join(cols) + " |",
           "| " + " | ".join("---" for _ in cols) + " |"]
    for r in rows:
        cells = []
        for c in cols:
            v = r.get(c, "")
            cells.append(f"{v:.6g}" if isinstance(v, float) else str(v))
        out.append("| " + " | ".join(cells) + " |")
    out.append("")
    return out


def to_markdown(doc: dict[str, Any], path: Path | None = None) -> str:
    """把 gather() 的结果排成 Markdown。"""
    L: list[str] = []
    a = L.append
    a("# 空间刚架分析报告")
    a("")
    a(f"生成时间：{doc['generated']}　·　模型单位制：{doc['units']['model']}"
      f"　·　报告单位：{doc['units']['report']}")
    a("")
    a("> 报告中每一个数字都来自求解工具的返回值，与界面、评测集走同一条计算链路，"
      "没有任何另行估算的成分。")
    a("")

    a("## 一、模型概况")
    a("")
    m = doc["model"]
    a(f"节点 {m['nodes']} 个，杆件 {m['members']} 根，支座 {m['supports']} 个"
      f"（其中全固接 {m['fully_fixed']} 个）。")
    if (m["analysis_nodes"] != m["nodes"]
            or m["analysis_elements"] != m["members"]):
        a(f"求解时自动编译为 {m['analysis_nodes']} 个分析节点、"
          f"{m['analysis_elements']} 个分析单元；上述数量仍按用户物理模型统计。")
    a("")
    a("**材料**")
    a("")
    L.extend(_table(m["materials"]))
    a("**截面**")
    a("")
    L.extend(_table(m["sections"]))
    a(f"荷载工况：{'、'.join(doc['cases']) or '（无）'}")
    a("")
    a(f"荷载组合：{'、'.join(doc['combos']) or '（无）'}")
    a("")

    a("## 二、求解与校验")
    a("")
    a("模型先过两级校验（JSON 结构 + 力学语义）才求解；刚度矩阵奇异会被主元检查"
      "拦下并定位到具体节点方向。下表是各工况的静力平衡校核。")
    a("")
    L.extend(_table(doc["checks"]))

    a("## 三、位移")
    a("")
    d = doc["displacement"]
    a(f"- 最大**节点**位移 {d['nodal_mm']:.4f} mm，位于节点 {d['node']}"
      f"（{d['coordinates'][0]:.3f}, {d['coordinates'][1]:.3f},"
      f" {d['coordinates'][2]:.3f}）")
    a(f"- 最大**杆件挠度** {d['deflection_mm']:.4f} mm，位于杆件 {d['at_member']}"
      f" 的 x = {d['at_x_m']:.3f} m 处")
    if d["length_over_deflection"]:
        a(f"- 该杆件长 {d['member_length_m']:.3f} m，杆长与挠度之比 "
          f"1/{d['length_over_deflection']:.0f}")
    a("")
    a("两点要说明。其一，节点位移**只在节点上取值**，单跨划一个单元时跨中根本没有"
      "节点，报出来的数会小一个量级；校核必须用杆件挠度——它由精确弯矩两次积分得到，"
      "与网格疏密无关。其二，上面那个比值的分母是**这一根杆件的长度，不是设计跨度**："
      "门式刚架的斜梁由两根杆组成，梁划成几个单元时差得更多。"
      "要对 L/400 这类限值，请按实际设计跨度另行换算。")
    a("")

    a("## 四、支座反力")
    a("")
    a(f"竖向反力合计 {doc['reactions']['total_vertical']:.3f} kN。")
    a("")
    L.extend(_table(doc["reactions"]["rows"]))

    if doc.get("envelope"):
        a("## 五、内力极值与控制组合")
        a("")
        a("包络逐点取各组合的上下界。同一根杆上跨中与支座常由不同组合控制，"
          "所以下表列出的控制组合是**对应那一点**的，不能推广到整根杆。")
        a("")
        L.extend(_table(doc["envelope"]))
        if doc.get("never_governs"):
            a(f"以下组合在任何杆件的任何一点都不控制：{'、'.join(doc['never_governs'])}。"
              "要么可以删掉，要么它们的系数或荷载写错了。")
            a("")

    if doc.get("modal"):
        a("## 六、自振特性")
        a("")
        rows = [{"阶": r["order"], "频率 (Hz)": r["frequency_Hz"],
                 "周期 (s)": r["period_s"]} for r in doc["modal"]["modes"]]
        L.extend(_table(rows))
        share = doc["modal"].get("effective_mass_ratio_xyz")
        if share:
            a(f"总质量 {doc['modal']['total_mass_kg']:.1f} kg；"
              f"前 {len(rows)} 阶的有效质量比 X/Y/Z = "
              f"{share[0]:.3f} / {share[1]:.3f} / {share[2]:.3f}。"
              "该比值接近 1 才说明取的阶数够。")
            a("")

    if doc.get("buckling"):
        a("## 七、稳定")
        a("")
        b = doc["buckling"]
        a(f"工况 {b['case']} 下的临界荷载因子 λ = **{b['critical_factor']:.3f}**"
          f"（前几阶：{'、'.join(f'{v:.3f}' for v in b['factors'])}）。")
        a("")
        a(f"最受压杆件是 {b['most_compressed_member']} 号，轴力 "
          f"{b['its_axial_kN']:.3f} kN。")
        a("")
        a("**这是线性特征值屈曲**，假定失稳前保持线弹性、变形小、轴力不随变形改变。"
          "真实结构有初始缺陷与残余应力，实际承载力低于此值——只能当上限，"
          "不能直接当承载力用。")
        a("")

    if doc["figures"]:
        a("## 八、图")
        a("")
        for label, fig in doc["figures"].items():
            a(f"**{label}**")
            a("")
            a(f"![{label}]({Path(fig['path']).name})")
            a("")

    if doc["skipped"]:
        a("## 附：未包含的内容")
        a("")
        a("下面这些没有出现在报告里，以及原因。列出来是为了让读的人知道"
          "它们是被明确跳过的，而不是忘了做。")
        a("")
        L.extend(_table([{"项目": k, "原因": v} for k, v in doc["skipped"].items()]))

    text = "\n".join(L)
    if path is not None:
        Path(path).write_text(text, encoding="utf-8")
    return text


# --------------------------------------------------------------- docx

_DOCX_HINT = ("需要 python-docx。在项目的虚拟环境里执行 "
              "`pip install python-docx` 即可；Markdown 版不需要它。")


def _set_cjk_font(run, name: str = "微软雅黑") -> None:
    """python-docx 默认只设西文字体，中文会回退到系统默认。

    Word 的字体是分东亚 / 西文两套的，只设 `run.font.name` 中文不生效——
    必须同时写 w:eastAsia。这是 python-docx 出中文文档最常见的坑。
    """
    from docx.oxml.ns import qn

    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        from docx.oxml import OxmlElement
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)


def _fmt(v: Any) -> str:
    """表格里的数怎么显示。

    先判量级再判整数：弹性模量 2.1e11 是个"整数"，直接 str() 会印出
    210000000000 这么一长串，表格立刻被撑变形。
    """
    if isinstance(v, bool) or v is None:
        return "" if v is None else str(v)
    if isinstance(v, (int, float)):
        x = float(v)
        if x != 0 and (abs(x) >= 1e5 or abs(x) < 1e-3):
            return f"{x:.4e}"
        if float(x).is_integer():
            return str(int(x))
        return f"{x:.4f}".rstrip("0").rstrip(".")
    return str(v)


def to_docx(doc: dict[str, Any], path: Path) -> Path:
    """产出 Word 文档。内容与 Markdown 版同源，只是排版不同。"""
    try:
        import docx as _docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:                      # pragma: no cover
        raise RuntimeError(f"{_DOCX_HINT}（{exc}）") from exc

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    d = _docx.Document()
    for style in ("Normal",):
        d.styles[style].font.size = Pt(10.5)

    def para(text: str = "", *, bold: bool = False, italic: bool = False,
             size: float = 10.5, grey: bool = False, center: bool = False):
        p = d.add_paragraph()
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
        run.font.size = Pt(size)
        if grey:
            run.font.color.rgb = RGBColor(0x6B, 0x6B, 0x6B)
        _set_cjk_font(run)
        return p

    def heading(text: str, level: int = 1):
        h = d.add_heading(text, level=level)
        for run in h.runs:
            _set_cjk_font(run)
        return h

    def table(rows: list[dict]) -> None:
        if not rows:
            para("（无）", italic=True)
            return
        cols = list(rows[0])
        t = d.add_table(rows=1, cols=len(cols))
        t.style = "Table Grid"
        for cell, name in zip(t.rows[0].cells, cols):
            run = cell.paragraphs[0].add_run(str(name))
            run.bold = True
            run.font.size = Pt(9)
            _set_cjk_font(run)
        for r in rows:
            cells = t.add_row().cells
            for cell, name in zip(cells, cols):
                run = cell.paragraphs[0].add_run(_fmt(r.get(name, "")))
                run.font.size = Pt(9)
                _set_cjk_font(run)
        para("", size=6)

    heading("空间刚架分析报告", 0)
    para(f"生成时间 {doc['generated']}　·　模型单位制 {doc['units']['model']}"
         f"　·　报告单位：{doc['units']['report']}", size=9, grey=True)
    para("报告中每一个数字都来自求解工具的返回值，与界面、评测集走同一条计算链路。",
         italic=True, size=9, grey=True)

    m = doc["model"]
    heading("一、模型概况", 1)
    para(f"节点 {m['nodes']} 个，杆件 {m['members']} 根，支座 {m['supports']} 个"
         f"（其中全固接 {m['fully_fixed']} 个）。")
    para("材料", bold=True)
    table(m["materials"])
    para("截面", bold=True)
    table(m["sections"])
    para("荷载工况：" + ("、".join(doc["cases"]) or "（无）"))
    para("荷载组合：" + ("、".join(doc["combos"]) or "（无）"))

    heading("二、求解与校验", 1)
    para("模型先过两级校验（JSON 结构 + 力学语义）才求解；刚度矩阵奇异会被主元"
         "检查拦下并定位到具体节点方向。")
    table(doc["checks"])

    dd = doc["displacement"]
    heading("三、位移", 1)
    para(f"最大节点位移 {dd['nodal_mm']:.4f} mm，位于节点 {dd['node']}。")
    ratio = (f"；该杆件长 {dd['member_length_m']:.3f} m，"
             f"杆长与挠度之比 1/{dd['length_over_deflection']:.0f}"
             if dd["length_over_deflection"] else "")
    para(f"最大杆件挠度 {dd['deflection_mm']:.4f} mm，位于杆件 {dd['at_member']}"
         f" 的 x = {dd['at_x_m']:.3f} m 处{ratio}。")
    para("两点要说明。其一，节点位移只在节点上取值，单跨划一个单元时跨中根本没有"
         "节点，报出来的数会小一个量级；校核必须用杆件挠度——它由精确弯矩两次积分"
         "得到，与网格疏密无关。其二，上面那个比值的分母是这一根杆件的长度，"
         "不是设计跨度：门式刚架的斜梁由两根杆组成，梁划成几个单元时差得更多。"
         "要对 L/400 这类限值，请按实际设计跨度另行换算。",
         size=9, grey=True)

    heading("四、支座反力", 1)
    para(f"竖向反力合计 {doc['reactions']['total_vertical']:.3f} kN。")
    table(doc["reactions"]["rows"])

    if doc.get("envelope"):
        heading("五、内力极值与控制组合", 1)
        para("包络逐点取各组合的上下界。同一根杆上跨中与支座常由不同组合控制，"
             "所以表中的控制组合是对应那一点的，不能推广到整根杆。")
        table(doc["envelope"])
        if doc.get("never_governs"):
            para("以下组合在任何杆件的任何一点都不控制："
                 + "、".join(doc["never_governs"])
                 + "。要么可以删掉，要么它们的系数或荷载写错了。",
                 italic=True, size=9, grey=True)

    if doc.get("modal"):
        heading("六、自振特性", 1)
        table([{"阶": r["order"], "频率 (Hz)": r["frequency_Hz"],
                "周期 (s)": r["period_s"]} for r in doc["modal"]["modes"]])
        share = doc["modal"].get("effective_mass_ratio_xyz")
        if share:
            para(f"总质量 {doc['modal']['total_mass_kg']:.1f} kg；有效质量比 "
                 f"X/Y/Z = {share[0]:.3f} / {share[1]:.3f} / {share[2]:.3f}。"
                 "该比值接近 1 才说明取的阶数够。", size=9, grey=True)

    if doc.get("buckling"):
        b = doc["buckling"]
        heading("七、稳定", 1)
        para(f"工况 {b['case']} 下的临界荷载因子 λ = {b['critical_factor']:.3f}"
             f"（前几阶：{'、'.join(f'{v:.3f}' for v in b['factors'])}）。"
             f"最受压杆件是 {b['most_compressed_member']} 号，轴力 "
             f"{b['its_axial_kN']:.3f} kN。")
        para("这是线性特征值屈曲，假定失稳前保持线弹性、变形小、轴力不随变形改变。"
             "真实结构有初始缺陷与残余应力，实际承载力低于此值——只能当上限，"
             "不能直接当承载力用。", italic=True, size=9, grey=True)

    if doc["figures"]:
        heading("八、图", 1)
        for label, fig in doc["figures"].items():
            para(label, bold=True)
            try:
                d.add_picture(fig["path"], width=Inches(5.6))
                d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception as exc:            # noqa: BLE001
                para(f"（图片读取失败：{exc}）", italic=True, size=9, grey=True)

    if doc["skipped"]:
        heading("附：未包含的内容", 1)
        para("列出来是为了让读的人知道它们是被明确跳过的，而不是忘了做。")
        table([{"项目": k, "原因": v} for k, v in doc["skipped"].items()])

    d.save(str(path))
    return path
